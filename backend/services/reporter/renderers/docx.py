from __future__ import annotations

import os
import tempfile
from typing import TYPE_CHECKING

from docx import Document
from docx.shared import Inches
from PIL import Image as PillowImage

from backend.services.reporter.renderers.sections import build_render_plan

if TYPE_CHECKING:
    from backend.services.reporter.config import ReporterConfig
    from backend.services.reporter.models import ParsedScan, ReproductionPackDraft


class DocxRenderer:
    def __init__(self, settings: ReporterConfig) -> None:
        self._settings = settings

    def generate(
        self,
        report_id: str,
        parsed_scan: ParsedScan,
        packs: list[ReproductionPackDraft],
    ) -> tuple[str, list[str]]:
        fd, output_path = tempfile.mkstemp(
            prefix=f"report-{report_id}-",
            suffix=".docx",
            dir=self._settings.temp_output_dir,
        )
        os.close(fd)

        plan = build_render_plan(
            parsed_scan=parsed_scan,
            packs=packs,
            include_raw_http_appendix=self._settings.include_raw_http_appendix,
        )

        doc = Document()
        partial_reasons: list[str] = []

        for line in plan.lines:
            if line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            else:
                doc.add_paragraph(line)

        if parsed_scan.include_evidence_screenshots:
            for entry in plan.evidence_images:
                doc.add_paragraph(f"Evidence image: {entry.finding_title}")
                try:
                    with PillowImage.open(entry.local_tmp_path) as image:
                        image.verify()
                    doc.add_picture(entry.local_tmp_path, width=Inches(5.5))
                except Exception:
                    doc.add_paragraph(
                        f"[Evidence screenshot unavailable - {entry.artifact_type}]"
                    )
                    partial_reasons.append("evidence_image_load_failed")

        doc.save(output_path)
        return output_path, _dedupe(partial_reasons)


def _dedupe(values: list[str]) -> list[str]:
    seen: set[str] = set()
    deduped: list[str] = []
    for value in values:
        if value in seen:
            continue
        seen.add(value)
        deduped.append(value)
    return deduped
